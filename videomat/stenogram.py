"""Stenogram: z adresu YouTube albo pliku do gotowego dokumentu z podziałem na mówców.

    from videomat.stenogram import Options, run
    wynik = run("https://www.youtube.com/watch?v=...", Options(title="Debata", speakers=3))

Przebieg: pobranie → ścieżka 16 kHz → podział długiego nagrania na kawałki z zakładką →
transkrypcja (faster-whisper) → diaryzacja (pyannote) → scalenie i grupowanie wypowiedzi →
eksport DOCX / Markdown / TXT / CSV / JSON + raport kontroli jakości.

Rozwiązania przejęte z pipeline'u NEXT100, bo sprawdziły się na sześciu godzinach konferencji:
cięcie kawałków w najcichszym miejscu (żeby nie urwać zdania), wyłączone
`condition_on_previous_text` (model przestaje halucynować pętle na oklaskach), rozpoznawanie
języka per segment po polskich znakach i słowach funkcyjnych, próg `avg_logprob < -0.8` na
fragmenty do odsłuchania, grupowanie wypowiedzi tego samego mówcy.
"""
from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from . import config, diarize as diarizer, ffmpeg, ingest, transcribe as asr

LOW_LOGPROB = -0.8
NO_SPEECH = 0.5
PL_MARKS = set("ąćęłńóśźżĄĆĘŁŃÓŚŹŻ")
PL_STOP = {"i", "w", "z", "na", "do", "się", "nie", "jest", "to", "że", "o", "jak", "ale", "co",
           "po", "dla", "być", "był", "od", "tak", "przez", "czy", "pan", "pani", "który", "bardzo"}
EN_STOP = {"the", "and", "of", "to", "in", "is", "that", "for", "it", "with", "as", "was", "on",
           "are", "this", "be", "have", "from", "we", "you", "they", "our"}


@dataclass
class Options:
    title: str = "Stenogram"
    subtitle: str = ""
    place: str = ""
    date: str = ""
    language: str | None = "pl"
    model: str = "large-v3"
    chunk_minutes: float = 15.0
    overlap_seconds: float = 15.0
    speakers: int | None = None
    min_speakers: int | None = None
    max_speakers: int | None = None
    speaker_names: dict[str, str] = field(default_factory=dict)
    voice_samples: dict[str, str] = field(default_factory=dict)
    diarize: bool = True
    section: tuple[float, float] | None = None
    group_gap: float = 3.0
    group_max: float = 75.0
    formats: tuple[str, ...] = ("docx", "md", "txt", "csv", "json")


def hms(seconds: float) -> str:
    seconds = max(0.0, float(seconds))
    h, rest = divmod(int(seconds), 3600)
    m, s = divmod(rest, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def slugify(text: str, limit: int = 50) -> str:
    import unicodedata
    plain = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    plain = re.sub(r"[^A-Za-z0-9]+", "_", plain).strip("_").lower()
    return (plain or "stenogram")[:limit]


# ------------------------------------------------------------------ dzielenie długiego nagrania
def quiet_point(wav: Path, around: float, window: float = 30.0) -> float:
    """Najcichsze miejsce w oknie wokół podanej sekundy — tam wypada ciąć, żeby nie urwać zdania."""
    try:
        import numpy as np
        import soundfile as sf
    except ImportError:
        return around
    info = sf.info(str(wav))
    rate = info.samplerate
    start = max(0.0, around - window)
    stop = min(info.duration, around + window)
    if stop - start < 1.0:
        return around
    data, _ = sf.read(str(wav), start=int(start * rate), stop=int(stop * rate), dtype="float32")
    if data.ndim > 1:
        data = data.mean(axis=1)
    frame = int(0.3 * rate)
    if frame <= 0 or len(data) < frame:
        return around
    frames = len(data) // frame
    energy = np.array([np.sqrt(np.mean(data[i * frame:(i + 1) * frame] ** 2)) for i in range(frames)])
    return start + float(energy.argmin()) * 0.3 + 0.15


def split_audio(wav: Path, out_dir: Path, chunk_minutes: float, overlap: float) -> list[dict]:
    """Kawałki z zakładką, cięte w ciszy. Krótkie nagranie zostaje w całości."""
    out_dir.mkdir(parents=True, exist_ok=True)
    total = ffmpeg.duration(wav)
    span = chunk_minutes * 60
    if total <= span * 1.2:
        return [{"path": str(wav), "offset": 0.0, "start": 0.0, "end": total}]
    chunks, cursor, index = [], 0.0, 0
    while cursor < total - 1.0:
        nominal = min(cursor + span, total)
        cut = total if nominal >= total else quiet_point(wav, nominal)
        piece = out_dir / f"chunk_{index:03d}.wav"
        if not piece.exists():
            ffmpeg.run(["-i", str(wav), "-ss", f"{cursor:.3f}", "-to", f"{min(cut + overlap, total):.3f}",
                        "-ac", "1", "-ar", "16000", str(piece)])
        chunks.append({"path": str(piece), "offset": round(cursor, 3),
                       "start": round(cursor, 3), "end": round(cut, 3)})
        cursor = cut
        index += 1
    return chunks


# ------------------------------------------------------------------ język i czyszczenie
def guess_language(text: str, previous: str = "pl") -> str:
    """Polski czy angielski — po znakach diakrytycznych, a przy ich braku po słowach funkcyjnych."""
    if any(ch in PL_MARKS for ch in text):
        return "pl"
    words = re.findall(r"[a-ząćęłńóśźż]+", text.lower())
    pl = sum(1 for w in words if w in PL_STOP)
    en = sum(1 for w in words if w in EN_STOP)
    if pl > en:
        return "pl"
    if en > pl:
        return "en"
    return previous


def drop_loops(segments: list[dict], repeats: int = 2, min_len: int = 12) -> list[dict]:
    """Usuwa pętle powtórzeń modelu — typowy artefakt na oklaskach i muzyce."""
    out: list[dict] = []
    streak, last = 0, None
    for segment in segments:
        text = segment.get("text", "").strip()
        if text and text == last and len(text) > min_len:
            streak += 1
            if streak >= repeats:
                continue
        else:
            streak = 0
        last = text
        out.append(segment)
    return out


def transcribe_long(wav: Path, work: Path, options: Options, on_progress=None) -> dict:
    """Transkrypcja z podziałem na kawałki i sklejeniem po osi czasu."""
    chunks = split_audio(wav, work / "chunks", options.chunk_minutes, options.overlap_seconds)
    segments: list[dict] = []
    uncertain: list[dict] = []
    last_end = -1.0
    previous_language = options.language or "pl"
    for index, chunk in enumerate(chunks):
        if on_progress:
            on_progress(0.15 + 0.45 * index / max(len(chunks), 1),
                        f"transkrypcja {index + 1}/{len(chunks)}")
        cache = work / "chunks" / f"chunk_{index:03d}.json"
        if cache.exists():
            part = json.loads(cache.read_text(encoding="utf-8"))
        else:
            part = asr.transcribe(chunk["path"], language=options.language, model_size=options.model)
            cache.write_text(json.dumps(part, ensure_ascii=False), encoding="utf-8")
        offset = chunk["offset"]
        for segment in part.get("segments", []):
            start = round(segment["start"] + offset, 2)
            end = round(segment["end"] + offset, 2)
            if start < last_end - 1.0:       # strefa zakładki — już to mamy
                continue
            language = guess_language(segment.get("text", ""), previous_language)
            previous_language = language
            words = [{**w, "start": round(w["start"] + offset, 2), "end": round(w["end"] + offset, 2)}
                     for w in segment.get("words", [])]
            low = (segment.get("avg_logprob", 0.0) < LOW_LOGPROB
                   or segment.get("no_speech_prob", 0.0) > NO_SPEECH)
            segments.append({"start": start, "end": end, "text": segment.get("text", "").strip(),
                             "words": words, "language": language, "low_confidence": bool(low),
                             "avg_logprob": segment.get("avg_logprob", 0.0),
                             "no_speech_prob": segment.get("no_speech_prob", 0.0)})
            last_end = max(last_end, end)
        uncertain += part.get("uncertain", [])
    return {"segments": drop_loops(segments), "uncertain": uncertain,
            "chunks": len(chunks), "model": options.model}


# ------------------------------------------------------------------ wypowiedzi
def group_utterances(segments: list[dict], gap: float = 3.0, cap: float = 75.0) -> list[dict]:
    """Skleja kolejne segmenty tego samego mówcy w wypowiedzi czytelne w dokumencie."""
    out: list[dict] = []
    for segment in segments:
        previous = out[-1] if out else None
        joinable = (previous and previous.get("speaker") == segment.get("speaker")
                    and segment["start"] - previous["end"] <= gap
                    and segment["end"] - previous["start"] <= cap
                    and previous.get("language") == segment.get("language"))
        if joinable:
            previous["end"] = segment["end"]
            previous["text"] = f'{previous["text"]} {segment["text"]}'.strip()
            previous["low_confidence"] = previous.get("low_confidence") or segment.get("low_confidence")
        else:
            out.append({k: segment.get(k) for k in
                        ("start", "end", "text", "speaker", "language", "low_confidence")})
    return out


def display_name(speaker: str | None, names: dict[str, str]) -> str:
    if not speaker:
        return "[MÓWCA NIEROZPOZNANY]"
    if speaker in names:
        return names[speaker]
    return f"[MÓWCA DO WERYFIKACJI – {speaker}]"


# ------------------------------------------------------------------ eksport
NAVY = "0B1A3F"
ACCENT = "C41E3A"
GREY = "555A66"


def export_docx(data: dict, path: Path, options: Options) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    navy = RGBColor(0x0B, 0x1A, 0x3F)
    accent = RGBColor(0xC4, 0x1E, 0x3A)
    grey = RGBColor(0x55, 0x5A, 0x66)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(options.title.upper())
    run.font.size = Pt(24)
    run.font.color.rgb = navy
    run.bold = True
    if options.subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub.add_run(options.subtitle)
        run.font.size = Pt(12)
        run.italic = True
        run.font.color.rgb = grey
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = label.add_run("S T E N O G R A M")
    run.font.size = Pt(13)
    run.font.color.rgb = accent
    run.bold = True

    meta_rows = [
        ("Źródło", data["source"].get("title") or data["source"].get("path", "")),
        ("Adres", data["source"].get("webpage_url") or "plik lokalny"),
        ("Długość", hms(data["duration"])),
        ("Data opracowania", options.date or datetime.now().strftime("%d.%m.%Y")),
        ("Miejsce", options.place or "—"),
        ("Mówcy", ", ".join(sorted({u["name"] for u in data["utterances"]})) or "—"),
        ("Model", f'{data["asr"]["model"]} · diaryzacja: {"tak" if data["diarized"] else "nie"}'),
        ("Status", data["qc"]["status"]),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for key, value in meta_rows:
        row = table.add_row().cells
        row[0].width = Cm(4.2)
        row[1].width = Cm(11.8)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), NAVY)
        row[0]._tc.get_or_add_tcPr().append(shade)
        run = row[0].paragraphs[0].add_run(key.upper())
        run.font.size = Pt(8.5)
        run.font.name = "Calibri"
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade2 = OxmlElement("w:shd")
        shade2.set(qn("w:fill"), "EEF1F7")
        row[1]._tc.get_or_add_tcPr().append(shade2)
        run = row[1].paragraphs[0].add_run(str(value))
        run.font.size = Pt(10.5)
        run.font.color.rgb = navy
    doc.add_paragraph()

    for utterance in data["utterances"]:
        paragraph = doc.add_paragraph()
        stamp = paragraph.add_run(f'[{hms(utterance["start"])}]  ')
        stamp.font.size = Pt(9.5)
        stamp.font.name = "Calibri"
        stamp.font.color.rgb = accent
        name = paragraph.add_run(utterance["name"])
        name.font.size = Pt(11.5)
        name.bold = True
        name.font.color.rgb = navy
        if utterance.get("language") and utterance["language"] != (options.language or "pl"):
            tag = paragraph.add_run(f'  [{utterance["language"].upper()}]')
            tag.font.size = Pt(8.5)
            tag.font.color.rgb = grey
        body = doc.add_paragraph()
        body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body.paragraph_format.line_spacing = 1.2
        body.paragraph_format.space_after = Pt(8)
        body.add_run(utterance["text"]).font.size = Pt(11)
        if utterance.get("low_confidence"):
            note = body.add_run("   [fragment o niskiej pewności — do odsłuchania]")
            note.font.size = Pt(8.5)
            note.italic = True
            note.font.name = "Calibri"
            note.font.color.rgb = RGBColor(0x8A, 0x8F, 0x99)

    flagged = [u for u in data["utterances"] if u.get("low_confidence")]
    if flagged:
        doc.add_page_break()
        heading = doc.add_paragraph().add_run("Miejsca wymagające weryfikacji")
        heading.bold = True
        heading.font.size = Pt(14)
        heading.font.color.rgb = navy
        check = doc.add_table(rows=1, cols=3)
        for i, title in enumerate(("Czas", "Mówca", "Fragment")):
            run = check.rows[0].cells[i].paragraphs[0].add_run(title)
            run.bold = True
            run.font.size = Pt(9)
        for utterance in flagged:
            cells = check.add_row().cells
            cells[0].text = hms(utterance["start"])
            cells[1].text = utterance["name"]
            cells[2].text = utterance["text"][:160]

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def export_markdown(data: dict, path: Path, options: Options) -> Path:
    lines = [f"# {options.title}", ""]
    if options.subtitle:
        lines += [f"*{options.subtitle}*", ""]
    source = data["source"]
    lines += [
        f"- Źródło: {source.get('title') or source.get('path')}",
        f"- Adres: {source.get('webpage_url') or 'plik lokalny'}",
        f"- Długość: {hms(data['duration'])}",
        f"- Model: {data['asr']['model']}, diaryzacja: {'tak' if data['diarized'] else 'nie'}",
        f"- Status kontroli: {data['qc']['status']}",
        "",
        "---",
        "",
    ]
    for utterance in data["utterances"]:
        flag = "  *(niska pewność)*" if utterance.get("low_confidence") else ""
        lines.append(f"**[{hms(utterance['start'])}] {utterance['name']}:**{flag}")
        lines.append("")
        lines.append(utterance["text"])
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def export_txt(data: dict, path: Path, options: Options) -> Path:
    lines = [options.title, "=" * len(options.title), ""]
    for utterance in data["utterances"]:
        lines.append(f"[{hms(utterance['start'])}] {utterance['name']}:")
        lines.append(utterance["text"])
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def export_csv(data: dict, path: Path) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["start", "end", "start_hms", "speaker", "name", "language",
                         "confidence", "text"])
        for utterance in data["utterances"]:
            writer.writerow([utterance["start"], utterance["end"], hms(utterance["start"]),
                             utterance.get("speaker") or "", utterance["name"],
                             utterance.get("language") or "",
                             "low" if utterance.get("low_confidence") else "ok",
                             utterance["text"]])
    return path


def build_qc(utterances: list[dict], segments: list[dict], turns: list[dict]) -> dict:
    flagged = [u for u in utterances if u.get("low_confidence")]
    share = len(flagged) / max(len(utterances), 1)
    status = ("Zweryfikowany" if not flagged
              else "Częściowo zweryfikowany" if share < 0.10 else "Wymaga kontroli")
    return {
        "segments": len(segments),
        "utterances": len(utterances),
        "low_confidence": len(flagged),
        "low_confidence_share": round(share, 3),
        "speakers": len({u.get("speaker") for u in utterances if u.get("speaker")}),
        "diarization_turns": len(turns),
        "status": status,
        "flagged": [{"start": hms(u["start"]), "name": u["name"], "text": u["text"][:200]}
                    for u in flagged[:80]],
    }


def export_qc(data: dict, path: Path, options: Options) -> Path:
    qc = data["qc"]
    lines = [
        f"# Kontrola jakości — {options.title}", "",
        f"- Segmentów rozpoznanych: {qc['segments']}",
        f"- Wypowiedzi po grupowaniu: {qc['utterances']}",
        f"- Rozpoznanych mówców: {qc['speakers']} (tur diaryzacji: {qc['diarization_turns']})",
        f"- Fragmentów o niskiej pewności: {qc['low_confidence']} "
        f"({qc['low_confidence_share'] * 100:.1f}%)",
        f"- Status: **{qc['status']}**", "",
        "Progi: `avg_logprob < -0.8` albo `no_speech_prob > 0.5` oznaczają fragment do odsłuchania.",
        "Nazwiska przypisane automatycznie wymagają potwierdzenia — bez mapowania mówcy zostają "
        "anonimowi, żeby nie podpisać wypowiedzi niewłaściwej osobie.", "",
    ]
    if qc["flagged"]:
        lines += ["## Do odsłuchania", "", "| Czas | Mówca | Fragment |", "|---|---|---|"]
        lines += [f"| {row['start']} | {row['name']} | {row['text'][:120].replace('|', '/')} |"
                  for row in qc["flagged"]]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


# ------------------------------------------------------------------ cały przebieg
def run(source: str | Path, options: Options | None = None, out_dir: Path | None = None,
        work_dir: Path | None = None, on_progress=None) -> dict:
    options = options or Options()
    config.ensure_dirs()
    step = on_progress or (lambda p, m: None)

    step(0.02, "pozyskuję materiał")
    work = Path(work_dir) if work_dir else config.WORK / "stenogram" / slugify(str(source))[:40]
    work.mkdir(parents=True, exist_ok=True)
    media = ingest.source_to_audio(source, work, audio_only=True, section=options.section)

    step(0.10, "przygotowuję ścieżkę dźwiękową")
    wav = ingest.to_wav16k(media["path"], work / "audio16k.wav")
    duration = ffmpeg.duration(wav)

    # Diaryzacja MUSI iść przed transkrypcją: faster-whisper ładuje własną wersję cuDNN,
    # po której pyannote nie potrafi już zainicjować GPU ("Could not load symbol cudnnGetLibConfig").
    turns: list[dict] = []
    diarized = False
    if options.diarize:
        try:
            step(0.12, "rozpoznaję mówców")
            turns = diarizer.diarize(wav, num_speakers=options.speakers,
                                     min_speakers=options.min_speakers,
                                     max_speakers=options.max_speakers,
                                     cache=work / "diarization.json")
            diarized = True
        except Exception as exc:
            step(0.12, f"diaryzacja pominięta: {str(exc)[:120]}")

    transcript = transcribe_long(wav, work, options, on_progress=step)

    names = dict(options.speaker_names)
    if diarized and options.voice_samples:
        try:
            step(0.72, "porównuję głosy z próbkami")
            names.update(diarizer.identify(wav, turns, options.voice_samples))
        except Exception as exc:
            step(0.72, f"rozpoznawanie osób pominięte: {str(exc)[:120]}")

    step(0.78, "składam wypowiedzi")
    segments = diarizer.attach(transcript, turns) if diarized else transcript["segments"]
    for segment in segments:
        segment.setdefault("language", options.language or "pl")
    utterances = group_utterances(segments, options.group_gap, options.group_max)
    for utterance in utterances:
        utterance["name"] = display_name(utterance.get("speaker"), names)

    data = {
        "source": media,
        "duration": duration,
        "asr": {"model": transcript.get("model"), "chunks": transcript.get("chunks")},
        "diarized": diarized,
        "speakers": diarizer.speaker_stats(utterances) if diarized else [],
        "names": names,
        "segments": segments,
        "utterances": utterances,
    }
    data["qc"] = build_qc(utterances, segments, turns)

    step(0.88, "zapisuję dokumenty")
    out_dir = Path(out_dir) if out_dir else config.OUT / "stenogramy"
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = slugify(options.title if options.title != "Stenogram" else
                   (media.get("title") or "stenogram"))
    written: dict[str, str] = {}
    if "json" in options.formats:
        target = out_dir / f"{stem}.json"
        target.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        written["json"] = str(target)
    if "csv" in options.formats:
        written["csv"] = str(export_csv(data, out_dir / f"{stem}.csv"))
    if "md" in options.formats:
        written["md"] = str(export_markdown(data, out_dir / f"{stem}.md", options))
    if "txt" in options.formats:
        written["txt"] = str(export_txt(data, out_dir / f"{stem}.txt", options))
    if "docx" in options.formats:
        try:
            written["docx"] = str(export_docx(data, out_dir / f"{stem}.docx", options))
        except ImportError:
            step(0.95, "brak python-docx — pomijam DOCX")
    written["qc"] = str(export_qc(data, out_dir / f"{stem}_QC.md", options))

    step(1.0, "gotowe")
    return {"files": written, "qc": data["qc"], "duration": duration,
            "utterances": len(utterances), "speakers": data["speakers"],
            "title": options.title, "source": media.get("webpage_url") or media.get("path")}
